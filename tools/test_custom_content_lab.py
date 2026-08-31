#!/usr/bin/env python3

import argparse
import json
import struct
import tempfile
import unittest
from pathlib import Path

import custom_content_lab as lab


def make_lev(path: Path) -> None:
    body_inst, body_nav, body_map = 0x220, 0x620, 0x700
    data = bytearray(4 + body_map + 4)
    def put(offset, value): struct.pack_into("<I", data, offset, value)
    put(0, body_map)
    models = (0x15, 0x23, 0x08, 0x5c, 0x93, 0x94, 0x95)
    put(4 + 0x0c, len(models)); put(4 + 0x10, body_inst); put(4 + 0x148, 35)
    for index, model in enumerate(models): put(4 + body_inst + index * 0x40 + 0x3c, model)
    for index in range(8): data[4 + 0x6c + index * 12] = index + 1
    put(4 + 0x188, body_nav)
    for index in range(3): put(4 + body_nav + index * 4, 0x680 + index * 0x10)
    path.write_bytes(data)


class LabTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="ctr-lab-test-")
        self.root = Path(self.temp.name)
        self.verifier = Path("/tmp/custom-content-verify")
        self.assertTrue(self.verifier.exists(), "compile /tmp/custom-content-verify first")

    def tearDown(self): self.temp.cleanup()

    def response(self):
        def media(mid, kind, modes):
            return {"id": mid, "type": kind, "version": "1.0.0", "modes": modes,
                    "file_size": 1, "crc32_hash": mid, "download_url": f"/file/{mid}"}
        return {"data": [
            {"id": 10, "name": "Fixture", "author": "Tester", "author_id": 2,
             "collection": "community", "collection_id": 2, "type": "race", "type_id": 1,
             "is_active": True, "downloads": [media(100, "lev", ["time_trial"]), media(101, "vrm", ["time_trial"])]},
            {"id": 11, "name": "Mismatch", "author": "Tester", "author_id": 2,
             "collection": "community", "collection_id": 2, "type": "race", "type_id": 1,
             "is_active": True, "downloads": [media(110, "lev", ["arcade"]), media(111, "vrm", [])]},
            {"id": 12, "name": "Hub", "author": "Tester", "author_id": 2,
             "collection": "community", "collection_id": 2, "type": "adventure_hub", "type_id": 3,
             "is_active": True, "downloads": []}]}

    def test_end_to_end_offline(self):
        catalogue = lab.catalogue_from_response(self.response())
        self.assertEqual(len(catalogue["tracks"]), 2)
        self.assertTrue(catalogue["tracks"][0]["versions"][0]["consistent"])
        self.assertEqual(len(catalogue["tracks"][1]["versions"]), 2)
        self.assertTrue(all(not row["consistent"] for row in catalogue["tracks"][1]["versions"]))
        lab.atomic_json(self.root / "catalogue.json", catalogue)
        lev, vrm = lab.cache_paths(self.root, 10, "1.0.0", 100, 101)
        lev.parent.mkdir(parents=True)
        make_lev(lev); vrm.write_bytes(b"vrm")
        lab.command_analyze(argparse.Namespace(root=self.root, verifier=self.verifier, offline=True))
        inv = lab.inventory(self.root)
        self.assertEqual(len(inv["entries"]), 3)
        good = inv["entries"][0]
        self.assertEqual(good["status"], "verified")
        self.assertEqual(good["verification"]["file_analysis"]["time_trial"]["result"], "Yes")
        self.assertEqual([row["status"] for row in inv["entries"][1:]], ["inconsistent", "inconsistent"])
        manifest = self.root / "session.json"
        lab.command_prepare_session(argparse.Namespace(root=self.root, variant_id=good["variant_id"],
                                    mode="time_trial", host=6, minutes=30, manifest=manifest))
        prepared = json.loads(manifest.read_text())
        self.assertEqual(prepared["variant_id"], good["variant_id"])
        self.assertEqual(prepared["host"], 6)
        executable = self.root / "ctr_native_ap"; executable.write_bytes(b"build")
        receipt = dict(prepared, executable_sha256=lab.sha256_file(executable),
                       git_commit="a" * 40, dirty=False, complete=True)
        receipt_path = self.root / "receipt.json"; lab.atomic_json(receipt_path, receipt)
        lab.command_collect_session(argparse.Namespace(root=self.root, manifest=manifest,
            receipt=receipt_path, executable=executable, reviewer="Tester", steam_attested=True,
            result="Yes", notes="fixture pass"))
        record = lab.verdicts(self.root)["verdicts"][good["variant_id"]]
        self.assertEqual(record["lev_sha256"], good["verification"]["lev_sha256"])
        self.assertEqual(record["modes"]["time_trial"]["launch_method"], "user-steam-attested")
        out_json, out_xlsx = self.root / "export.json", self.root / "export.xlsx"
        lab.command_export(argparse.Namespace(root=self.root, json=out_json, xlsx=out_xlsx))
        exported = json.loads(out_json.read_text())
        self.assertEqual(exported["entries"][0]["time_trial In-game test"], "Yes")
        self.assertTrue(out_xlsx.exists() and out_xlsx.stat().st_size > 0)

    def test_untrusted_boundaries(self):
        with self.assertRaises(ValueError): lab.positive_int(True, "id")
        with self.assertRaises(ValueError): lab.saphi_url("https://evil.invalid/file")
        with self.assertRaises(ValueError): lab.saphi_url("//evil.invalid/file")
        lev, vrm = lab.cache_paths(self.root, 1, "../../escape", 2, 3)
        self.assertEqual(lev.parent, self.root / "media" / "1")
        self.assertEqual(vrm.parent, self.root / "media" / "1")

    def test_formula_text_is_neutralized(self):
        rows = [{"name": " =1+1", "author": "+cmd", "status": "verified"}]
        original = lab.rows_for_export
        lab.rows_for_export = lambda root: rows
        try:
            out_json, out_xlsx = self.root / "formula.json", self.root / "formula.xlsx"
            lab.command_export(argparse.Namespace(root=self.root, json=out_json, xlsx=out_xlsx))
            from openpyxl import load_workbook
            workbook = load_workbook(out_xlsx, data_only=False)
            values = [cell for row in workbook.active.iter_rows(min_row=2) for cell in row]
            self.assertTrue(all(cell.data_type == "s" for cell in values))
            self.assertTrue(values[0].value.startswith("'"))
        finally:
            lab.rows_for_export = original


if __name__ == "__main__": unittest.main()
