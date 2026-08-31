#!/usr/bin/env python3
"""Project Saphi v3 catalogue, cache, verification ledger and exporters.

This is the orchestration layer around the native verifier. It never mutates a
downloaded creator file. State changes are atomic and every runtime verdict is
bound to the exact LEV and VRM SHA-256 pair reported by the native core.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import shutil
import secrets
import subprocess
import sys
import tempfile
import time
import zlib
import urllib.error
import urllib.parse
import urllib.request
from email.utils import parsedate_to_datetime
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

API_ROOT = "https://www.projectsaphi.com"
TRACKS_URL = f"{API_ROOT}/api/v3/tracks?per_page=100&include_standards=0&include_events=0&include_downloads=1"
MODES = ("time_trial", "relic_race", "ctr_challenge", "arcade", "crystal_challenge")
RESULTS = ("Yes", "No", "Not tested")
STATIC_RESULTS = ("Detected", "Not detected", "Indeterminate", "Error")
SCHEMA_VERSION = 2
MAX_JSON_BYTES = 64 * 1024 * 1024
MAX_TEXT = 512
MAX_URL = 2048
MAX_PAGE_BYTES = 16 * 1024 * 1024
MAX_PAGES = 1000
MAX_TRACKS = 100000
MAX_VERSIONS = 100
MAX_LEDGER_RECORDS = 500000


def check_json_depth(value: Any, label: str, maximum: int = 64) -> None:
    def visit(node: Any, depth: int) -> None:
        if depth > maximum: raise ValueError(f"JSON nesting exceeds limit: {label}")
        if isinstance(node, dict):
            for child in node.values(): visit(child, depth + 1)
        elif isinstance(node, list):
            for child in node: visit(child, depth + 1)
    visit(value, 0)


def atomic_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix=f".{path.name}.", suffix=".tmp", dir=path.parent)
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as stream:
            json.dump(value, stream, indent=2, ensure_ascii=False, sort_keys=True)
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
        directory = os.open(path.parent, os.O_RDONLY)
        try: os.fsync(directory)
        finally: os.close(directory)
    except BaseException:
        try:
            os.unlink(temporary)
        except FileNotFoundError:
            pass
        raise


def load_json(path: Path, default: Any = None) -> Any:
    if not path.exists():
        return default
    if path.stat().st_size > MAX_JSON_BYTES:
        raise ValueError(f"JSON exceeds {MAX_JSON_BYTES} byte limit: {path}")
    with path.open("r", encoding="utf-8") as stream:
        value = json.load(stream)
    check_json_depth(value, str(path))
    return value


def fetch_json(url: str) -> Any:
    request = urllib.request.Request(url, headers={"User-Agent": "CTR-AP-Verification-Lab/0.1"})
    with urllib.request.urlopen(request, timeout=30) as response:
        raw = response.read(MAX_PAGE_BYTES + 1)
        if len(raw) > MAX_PAGE_BYTES: raise ValueError("Saphi response exceeds page limit")
        value = json.loads(raw)
        check_json_depth(value, url)
        return value


def positive_int(value: Any, field: str) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise ValueError(f"{field} must be a positive integer")
    return value


def bounded_text(value: Any, field: str) -> str:
    if not isinstance(value, str) or len(value) > MAX_TEXT or "\x00" in value:
        raise ValueError(f"{field} must be a string of at most {MAX_TEXT} characters")
    return value


def normalized_crc32(value: Any) -> str:
    if isinstance(value, bool): raise ValueError("CRC32 must be hexadecimal text or integer")
    if isinstance(value, int):
        if value < 0 or value > 0xffffffff: raise ValueError("CRC32 integer out of range")
        return f"{value:08x}"
    if isinstance(value, str):
        text = value.lower().removeprefix("0x")
        if len(text) == 8 and all(ch in "0123456789abcdef" for ch in text): return text
    raise ValueError("CRC32 must contain exactly eight hexadecimal digits")


def saphi_url(value: Any) -> str:
    value = bounded_text(value, "download_url")
    if len(value) > MAX_URL: raise ValueError("download_url exceeds limit")
    resolved = urllib.parse.urljoin(API_ROOT + "/", value)
    parsed = urllib.parse.urlsplit(resolved)
    expected = urllib.parse.urlsplit(API_ROOT)
    if parsed.scheme != "https" or parsed.hostname != expected.hostname or parsed.port not in (None, 443) or parsed.username or parsed.password:
        raise ValueError("download_url must stay on the configured Project Saphi HTTPS origin")
    return resolved


def media_versions(track: dict[str, Any]) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, list[dict[str, Any]]]] = {}
    downloads = track.get("downloads", [])
    if not isinstance(downloads, list) or len(downloads) > MAX_VERSIONS * 2:
        raise ValueError("downloads must be a bounded list")
    seen_media: set[int] = set()
    for media in downloads:
        if not isinstance(media, dict): raise ValueError("download must be an object")
        if media.get("type") not in {"lev", "vrm"} or not media.get("download_url"):
            continue
        mid = positive_int(media.get("id"), "media.id")
        if mid in seen_media: raise ValueError(f"duplicate media id {mid}")
        seen_media.add(mid)
        version = bounded_text(media.get("version"), "media.version")
        if not version: raise ValueError("media.version must not be empty")
        modes = media.get("modes", [])
        if not isinstance(modes, list) or any(mode not in MODES for mode in modes):
            raise ValueError("media.modes contains an unsupported tag")
        normalized = dict(media)
        normalized.update({"id": mid, "file_size": positive_int(media.get("file_size"), "media.file_size"),
                           "crc32_hash": normalized_crc32(media.get("crc32_hash")),
                           "download_url": saphi_url(media["download_url"]), "version": version,
                           "modes": list(dict.fromkeys(modes))})
        grouped.setdefault(version, {"lev": [], "vrm": []})[media["type"]].append(normalized)
    if len(grouped) > MAX_VERSIONS: raise ValueError("track has too many versions")
    rows = []
    for version, pair in sorted(grouped.items()):
        signatures = sorted({tuple(sorted(row.get("modes", []))) for row in pair["lev"] + pair["vrm"]})
        for signature in signatures:
            levs = [row for row in pair["lev"] if tuple(sorted(row.get("modes", []))) == signature]
            vrms = [row for row in pair["vrm"] if tuple(sorted(row.get("modes", []))) == signature]
            reasons = []
            if len(levs) != 1: reasons.append(f"expected one LEV for modes {signature!r}, found {len(levs)}")
            if len(vrms) != 1: reasons.append(f"expected one VRM for modes {signature!r}, found {len(vrms)}")
            rows.append({"version": version, "lev": levs[0] if len(levs) == 1 else None,
                         "vrm": vrms[0] if len(vrms) == 1 else None,
                         "saphi_modes": list(signature), "consistent": not reasons,
                         "inconsistency": "; ".join(reasons)})
    return rows


def catalogue_from_response(response: dict[str, Any]) -> dict[str, Any]:
    if not isinstance(response, dict) or not isinstance(response.get("data"), list):
        raise ValueError("Saphi response must contain a data list")
    tracks = []
    seen_tracks: set[int] = set()
    seen_media_global: set[int] = set()
    for track in response.get("data", []):
        if not isinstance(track, dict): raise ValueError("track must be an object")
        if track.get("type") not in {"race", "battle"}:
            continue
        tid = positive_int(track.get("id"), "track.id")
        if not isinstance(track.get("is_active"), bool): raise ValueError("track.is_active must be boolean")
        if tid in seen_tracks: raise ValueError(f"duplicate track id {tid}")
        seen_tracks.add(tid)
        versions = media_versions(track)
        for version in versions:
            for role in ("lev", "vrm"):
                media = version.get(role)
                if not media: continue
                if media["id"] in seen_media_global: raise ValueError(f"duplicate global media id {media['id']}")
                seen_media_global.add(media["id"])
        tracks.append({"id": tid, "name": bounded_text(track.get("name"), "track.name"),
                       "author_id": positive_int(track.get("author_id"), "track.author_id"),
                       "author": bounded_text(track.get("author"), "track.author"),
                       "collection_id": positive_int(track.get("collection_id"), "track.collection_id"),
                       "collection": bounded_text(track.get("collection"), "track.collection"),
                       "type_id": positive_int(track.get("type_id"), "track.type_id"),
                       "type": track["type"], "is_active": track.get("is_active"),
                       "versions": versions})
        if len(tracks) > MAX_TRACKS: raise ValueError("normalized track limit exceeded")
    return {"schema_version": SCHEMA_VERSION,
            "source": TRACKS_URL,
            "fetched_at": datetime.now(timezone.utc).isoformat(),
            "tracks": tracks}


_last_download_request = 0.0


def file_crc32(path: Path) -> str:
    value = 0
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(128 * 1024), b""):
            value = zlib.crc32(chunk, value)
    return f"{value & 0xffffffff:08x}"


def quarantine(path: Path, reason: str) -> None:
    if path.exists():
        stamp = int(time.time_ns())
        path.replace(path.with_name(f"{path.name}.quarantine-{stamp}-{reason}"))


def valid_cached_media(media: dict[str, Any], target: Path) -> bool:
    try:
        return (target.is_file() and not target.is_symlink() and
                target.stat().st_size == media["file_size"] and
                file_crc32(target) == media["crc32_hash"])
    except OSError:
        return False


def download_media(media: dict[str, Any], target: Path) -> None:
    expected = positive_int(media.get("file_size"), "media.file_size")
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists() and valid_cached_media(media, target):
        return
    if target.exists(): quarantine(target, "integrity")
    partial = target.with_suffix(target.suffix + ".part")
    start = partial.stat().st_size if partial.exists() else 0
    if start > expected:
        quarantine(partial, "oversized")
        start = 0
    headers = {"User-Agent": "CTR-AP-Verification-Lab/0.1"}
    if start:
        headers["Range"] = f"bytes={start}-"
    url = saphi_url(media["download_url"])
    global _last_download_request
    for attempt in range(6):
        delay = 0.55 - (time.monotonic() - _last_download_request)
        if delay > 0: time.sleep(delay)
        try:
            _last_download_request = time.monotonic()
            request = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(request, timeout=60) as response:
                status = getattr(response, "status", None)
                resumed = start > 0 and status == 206
                if resumed:
                    content_range = response.headers.get("Content-Range", "")
                    match = re.fullmatch(r"bytes (\d+)-(\d+)/(\d+)", content_range)
                    if not match or int(match.group(1)) != start or int(match.group(3)) != expected or int(match.group(2)) < start:
                        raise IOError("invalid Content-Range for resumed download")
                elif start and status == 200:
                    start = 0
                elif start:
                    raise IOError("range request did not return 206 or restartable 200")
                mode = "ab" if resumed else "wb"
                with partial.open(mode) as stream:
                    remaining = expected - start
                    while remaining:
                        chunk = response.read(min(128 * 1024, remaining))
                        if not chunk: raise IOError("early EOF while downloading media")
                        stream.write(chunk)
                        remaining -= len(chunk)
                    if response.read(1): raise IOError("download exceeded Saphi file_size")
                    stream.flush()
                    os.fsync(stream.fileno())
            break
        except urllib.error.HTTPError as exc:
            if exc.code == 416 and partial.exists() and valid_cached_media(media, partial):
                break
            if exc.code != 429 or attempt == 5: raise
            retry_after = exc.headers.get("Retry-After")
            try:
                wait = float(retry_after)
                if wait < 0: raise ValueError
            except (TypeError, ValueError):
                try:
                    when = parsedate_to_datetime(retry_after)
                    if when.tzinfo is None: when = when.replace(tzinfo=timezone.utc)
                    wait = max(0.0, (when - datetime.now(timezone.utc)).total_seconds())
                except (TypeError, ValueError, OverflowError):
                    wait = min(60.0, (2.0 ** attempt) + secrets.randbelow(1000) / 1000.0)
            wait = min(60.0, wait)
            time.sleep(wait)
    if partial.stat().st_size != expected:
        raise IOError(f"incomplete download: got {partial.stat().st_size}, expected {expected}")
    if file_crc32(partial) != normalized_crc32(media.get("crc32_hash")):
        quarantine(partial, "crc32")
        raise IOError("downloaded media CRC32 does not match Saphi")
    os.replace(partial, target)
    directory = os.open(target.parent, os.O_RDONLY)
    try: os.fsync(directory)
    finally: os.close(directory)


def cache_paths(root: Path, track_id: int, version: str, lev_id: int, vrm_id: int) -> tuple[Path, Path]:
    del version  # public semantic versions are metadata, never path components
    root = root.resolve()
    base = root / "media" / str(positive_int(track_id, "track_id"))
    lev = base / f"lev-{positive_int(lev_id, 'lev_media_id')}.lev"
    vrm = base / f"vrm-{positive_int(vrm_id, 'vrm_media_id')}.vrm"
    for candidate in (lev, vrm):
        if root not in candidate.resolve(strict=False).parents:
            raise ValueError("cache path escapes configured root")
        cursor = candidate.parent
        while cursor != root:
            if cursor.exists() and cursor.is_symlink(): raise ValueError("symlinked cache path refused")
            cursor = cursor.parent
    return lev, vrm


def identity_hash(parts: list[str]) -> str:
    if any(not isinstance(part, str) or "\x00" in part for part in parts):
        raise ValueError("identity fields must be canonical NUL-free strings")
    return hashlib.sha256("\0".join(parts).encode("utf-8")).hexdigest()


def content_id(report: dict[str, Any]) -> str:
    return identity_hash([report["lev_sha256"], report["vrm_sha256"]])


def variant_id(track_id: int, version: str, lev_id: int, vrm_id: int) -> str:
    return identity_hash([str(positive_int(track_id, "track_id")), bounded_text(version, "version"),
                          str(positive_int(lev_id, "lev_media_id")), str(positive_int(vrm_id, "vrm_media_id"))])


def run_verifier(executable: Path, lev: Path, vrm: Path) -> dict[str, Any]:
    result = subprocess.run([str(executable), str(lev), str(vrm)], text=True,
                            stdout=subprocess.PIPE, stderr=subprocess.PIPE, check=False)
    if result.returncode:
        raise RuntimeError(result.stderr.strip() or f"verifier exited {result.returncode}")
    return json.loads(result.stdout)


def inventory(root: Path) -> dict[str, Any]:
    data = load_json(root / "inventory.json", {"schema_version": SCHEMA_VERSION, "entries": []})
    if not isinstance(data, dict) or not isinstance(data.get("entries"), list):
        raise ValueError("inventory must contain an entries list")
    for entry in data["entries"]:
        if not isinstance(entry, dict): raise ValueError("inventory entry must be an object")
        if entry.get("variant_id"):
            expected = variant_id(entry["track_id"], entry["version"], entry["lev_media_id"], entry["vrm_media_id"])
            if entry["variant_id"] != expected: raise ValueError("serialized variant_id mismatch")
        if entry.get("content_id"):
            expected = content_id(entry.get("verification", {}))
            if entry["content_id"] != expected: raise ValueError("serialized content_id mismatch")
    return data


def verdicts(root: Path) -> dict[str, Any]:
    data = load_json(root / "verdicts.json", {"schema_version": SCHEMA_VERSION, "verdicts": {}})
    if not isinstance(data, dict) or not isinstance(data.get("verdicts"), dict):
        raise ValueError("verdict ledger must contain a verdicts object")
    if len(data["verdicts"]) > MAX_LEDGER_RECORDS: raise ValueError("verdict ledger record limit exceeded")
    for key, record in data["verdicts"].items():
        if not isinstance(record, dict) or record.get("variant_id") != key:
            raise ValueError("verdict ledger key does not match serialized variant_id")
        expected = variant_id(record["track_id"], record["version"], record["lev_media_id"], record["vrm_media_id"])
        if key != expected: raise ValueError("verdict ledger variant_id mismatch")
        expected_content = identity_hash([record["lev_sha256"], record["vrm_sha256"]])
        if record.get("content_id") != expected_content: raise ValueError("verdict ledger content_id mismatch")
    return data


def command_sync(args: argparse.Namespace) -> None:
    if args.catalogue_file:
        response = load_json(args.catalogue_file)
    else:
        response = None
        combined: list[Any] = []
        url = TRACKS_URL
        seen_urls: set[str] = set()
        for _ in range(MAX_PAGES):
            if url in seen_urls: raise ValueError("Saphi pagination loop detected")
            seen_urls.add(url)
            page = fetch_json(url)
            if not isinstance(page, dict) or not isinstance(page.get("data"), list):
                raise ValueError("invalid Saphi page")
            combined.extend(page["data"])
            if len(combined) > MAX_TRACKS: raise ValueError("Saphi total track limit exceeded")
            next_url = page.get("links", {}).get("next") if isinstance(page.get("links"), dict) else None
            if next_url:
                url = saphi_url(next_url)
                continue
            meta = page.get("meta", {}) if isinstance(page.get("meta"), dict) else {}
            current, last = meta.get("current_page"), meta.get("last_page")
            if current is not None or last is not None:
                if (isinstance(current, bool) or isinstance(last, bool) or not isinstance(current, int) or
                    not isinstance(last, int) or current < 1 or last < current):
                    raise ValueError("inconsistent Saphi page metadata")
            if isinstance(current, int) and isinstance(last, int) and current < last:
                parsed = urllib.parse.urlsplit(url)
                query = urllib.parse.parse_qs(parsed.query)
                query["page"] = [str(current + 1)]
                url = urllib.parse.urlunsplit(parsed._replace(query=urllib.parse.urlencode(query, doseq=True)))
                continue
            response = dict(page, data=combined)
            break
        else: raise ValueError("Saphi page limit exceeded")
    catalogue = catalogue_from_response(response)
    atomic_json(args.root / "catalogue.json", catalogue)
    print(f"catalogued {len(catalogue['tracks'])} race/battle entries")


def command_analyze(args: argparse.Namespace) -> None:
    catalogue = load_json(args.root / "catalogue.json")
    if not catalogue: raise SystemExit("run sync first")
    previous = inventory(args.root)
    existing = {entry["variant_id"]: entry for entry in previous["entries"] if entry.get("variant_id")}
    entries: list[dict[str, Any]] = []
    desired_ids = {variant_id(track["id"], version["version"], version["lev"]["id"], version["vrm"]["id"])
                   for track in catalogue["tracks"] for version in track["versions"]
                   if version.get("consistent") and version.get("lev") and version.get("vrm")}
    processed_ids: set[str] = set()
    scan_id = secrets.token_hex(16)
    catalogue_digest = hashlib.sha256(json.dumps(catalogue, sort_keys=True, separators=(",", ":")).encode()).hexdigest()

    def checkpoint(last: str | None) -> None:
        preserved = [entry for key, entry in existing.items()
                     if key in desired_ids and key not in processed_ids]
        atomic_json(args.root / "inventory.json", {"schema_version": SCHEMA_VERSION,
                    "scan_id": scan_id, "catalogue_digest": catalogue_digest,
                    "last_completed_identity": last, "entries": entries + preserved})

    for track in catalogue["tracks"]:
        for version in track["versions"]:
            base = {"track_id": track["id"], "name": track["name"], "author": track["author"],
                    "track_type": track["type"], "collection": track["collection"],
                    "version": version["version"], "saphi_modes": version["saphi_modes"],
                    "consistent": version["consistent"], "notes": version["inconsistency"]}
            if not version["consistent"]:
                base["status"] = "inconsistent"
                entries.append(base)
                checkpoint(None)
                continue
            lev_media, vrm_media = version["lev"], version["vrm"]
            base.update({"lev_media_id": lev_media["id"], "vrm_media_id": vrm_media["id"],
                         "lev_crc32": lev_media.get("crc32_hash"), "vrm_crc32": vrm_media.get("crc32_hash")})
            lev, vrm = cache_paths(args.root, track["id"], version["version"], lev_media["id"], vrm_media["id"])
            vid = variant_id(track["id"], version["version"], lev_media["id"], vrm_media["id"])
            prior = existing.get(vid)
            metadata_unchanged = prior is not None and all(prior.get(key) == base.get(key) for key in
                ("track_id", "name", "author", "track_type", "collection", "version",
                 "saphi_modes", "consistent", "lev_media_id", "vrm_media_id", "lev_crc32", "vrm_crc32"))
            if metadata_unchanged and prior.get("status") == "verified" and \
               valid_cached_media(lev_media, lev) and valid_cached_media(vrm_media, vrm) and \
               sha256_file(lev) == prior.get("verification", {}).get("lev_sha256") and \
               sha256_file(vrm) == prior.get("verification", {}).get("vrm_sha256"):
                entries.append(prior)
                processed_ids.add(vid)
                checkpoint(vid)
                continue
            try:
                if not args.offline:
                    download_media(lev_media, lev)
                    download_media(vrm_media, vrm)
                report = run_verifier(args.verifier, lev, vrm)
                base.update({"variant_id": vid, "content_id": content_id(report), "status": "verified", "lev_path": str(lev.relative_to(args.root)),
                             "vrm_path": str(vrm.relative_to(args.root)), "verification": report})
                if prior and prior.get("notes") and not base["notes"]:
                    base["notes"] = prior["notes"]
            except Exception as exc:  # retained as evidence; one bad track must not erase the sweep
                base.update({"status": "error", "notes": str(exc)})
            entries.append(base)
            processed_ids.add(vid)
            checkpoint(vid)
    print(f"analysed {len(entries)} exact versions")


def command_queue(args: argparse.Namespace) -> None:
    if args.cursor < 0: raise SystemExit("cursor must be nonnegative")
    ledger = verdicts(args.root)["verdicts"]
    rows = []
    for entry in inventory(args.root)["entries"]:
        if entry.get("status") != "verified": continue
        if args.author and entry.get("author") != args.author: continue
        if args.track_type and entry.get("track_type") != args.track_type: continue
        if args.collection and entry.get("collection") != args.collection: continue
        if args.mode and args.mode not in entry.get("saphi_modes", []): continue
        values = ledger.get(entry["variant_id"], {}).get("modes", {})
        pending = [mode for mode in MODES if values.get(mode, {}).get("result", "Not tested") == "Not tested"]
        if args.pending and not pending: continue
        report = entry.get("verification", {}).get("file_analysis", {})
        differences = [mode for mode in MODES if
                       (report.get(mode, {}).get("result") == "Detected") != (mode in entry.get("saphi_modes", []))]
        if args.difference and not differences: continue
        rows.append(dict(entry, pending=pending, differences=differences))

    def version_key(value: Any) -> tuple[Any, ...]:
        text = str(value)
        match = re.fullmatch(r"v?(\d+)\.(\d+)\.(\d+)(?:-([0-9A-Za-z.-]+))?(?:\+[0-9A-Za-z.-]+)?", text)
        if not match: return (1, text.casefold(), text)
        prerelease = match.group(4)
        identifiers = tuple((0, int(part)) if part.isdigit() else (1, part.casefold())
                            for part in prerelease.split(".")) if prerelease else ()
        return (0, int(match.group(1)), int(match.group(2)), int(match.group(3)),
                0 if prerelease else 1, identifiers)
    rows.sort(key=lambda row: (row["track_id"], version_key(row["version"]),
                               row["lev_media_id"], row["vrm_media_id"], row["variant_id"]))
    if args.newest_only:
        newest: dict[int, dict[str, Any]] = {}
        for row in rows:
            current = newest.get(row["track_id"])
            candidate_key, current_key = version_key(row["version"]), version_key(current["version"]) if current else None
            if current is None or (candidate_key[0] == 0 and current_key[0] != 0) or \
               (candidate_key[0] == current_key[0] and candidate_key > current_key):
                newest[row["track_id"]] = row
        rows = sorted(newest.values(), key=lambda row: (row["track_id"], row["variant_id"]))
    start = args.cursor
    page = rows[start:start + args.page_size]
    if args.json:
        print(json.dumps({"rows": page, "next_cursor": start + len(page) if start + len(page) < len(rows) else None}, sort_keys=True))
    else:
        for entry in page:
            print(f"{entry['variant_id'][:12]}  Saphi {entry['track_id']} media {entry['lev_media_id']}/{entry['vrm_media_id']} {entry['name']} {entry['version']}  {','.join(entry['pending'])}")


def verified_entry(root: Path, key: str) -> dict[str, Any]:
    entry = next((row for row in inventory(root)["entries"]
                  if row.get("variant_id") == key and row.get("status") == "verified"), None)
    if not entry: raise SystemExit("unknown or unverified exact-version key")
    return entry


def command_prepare_session(args: argparse.Namespace) -> None:
    entry = verified_entry(args.root, args.variant_id)
    if entry["track_type"] != "race": raise SystemExit("battle runtime review is Unsupported")
    if args.mode == "crystal_challenge": raise SystemExit("crystal_challenge battle review is Unsupported")
    if args.host < 0 or args.host > 17: raise SystemExit("host must be an arcade level ID from 0 through 17")
    if args.minutes < 1 or args.minutes > 60: raise SystemExit("session lifetime must be from 1 through 60 minutes")
    report = entry["verification"]
    lev_path = (args.root / entry["lev_path"]).resolve()
    vrm_path = (args.root / entry["vrm_path"]).resolve()
    if (not lev_path.is_file() or lev_path.is_symlink() or lev_path.stat().st_size != report["lev_bytes"] or
        file_crc32(lev_path) != entry["lev_crc32"] or sha256_file(lev_path) != report["lev_sha256"] or
        not vrm_path.is_file() or vrm_path.is_symlink() or vrm_path.stat().st_size != report["vrm_bytes"] or
        file_crc32(vrm_path) != entry["vrm_crc32"] or sha256_file(vrm_path) != report["vrm_sha256"]):
        raise SystemExit("cached pair no longer matches inventory; analyze again")
    session_dir = args.root / "session"
    manifest_path = session_dir / "manifest.json"
    armed_path = session_dir / "armed.json"
    receipt_path = session_dir / "receipt.json"
    if manifest_path.exists() or armed_path.exists() or receipt_path.exists():
        raise SystemExit("a prepared or uncollected session already exists")
    now = datetime.now(timezone.utc)
    expiry = now + timedelta(minutes=args.minutes)
    manifest = {"schema_version": SCHEMA_VERSION, "nonce": secrets.token_hex(32),
                "variant_id": entry["variant_id"], "content_id": entry["content_id"],
                "track_id": entry["track_id"], "version": entry["version"],
                "lev_media_id": entry["lev_media_id"], "vrm_media_id": entry["vrm_media_id"],
                "lev_path": str(lev_path), "vrm_path": str(vrm_path),
                "lev_sha256": report["lev_sha256"], "vrm_sha256": report["vrm_sha256"],
                "lev_bytes": report["lev_bytes"], "vrm_bytes": report["vrm_bytes"],
                "host": args.host, "mode": args.mode,
                "created_at": now.isoformat(), "expires_at": expiry.isoformat(),
                "created_unix": int(now.timestamp()), "expires_unix": int(expiry.timestamp())}
    atomic_json(manifest_path, manifest)
    print(f"Prepared one-use Steam review session {manifest['nonce']} at {manifest_path}.")
    print("Start the configured development build through Steam. The lab did not launch CTR.")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(128 * 1024), b""): digest.update(chunk)
    return digest.hexdigest()


def command_collect_session(args: argparse.Namespace) -> None:
    if not args.steam_attested: raise SystemExit("reviewer Steam-launch attestation is required")
    manifest_path = args.root / "session" / "armed.json"
    receipt_path = args.root / "session" / "receipt.json"
    manifest = load_json(manifest_path)
    receipt = load_json(receipt_path)
    if not isinstance(manifest, dict) or not isinstance(receipt, dict): raise SystemExit("session manifest and receipt are required")
    expires = datetime.fromisoformat(manifest["expires_at"])
    if datetime.now(timezone.utc) > expires: raise SystemExit("session manifest expired")
    fields = ("nonce", "variant_id", "content_id", "track_id", "version", "lev_media_id", "vrm_media_id",
              "lev_sha256", "vrm_sha256", "lev_bytes", "vrm_bytes", "host", "mode")
    if any(receipt.get(field) != manifest.get(field) for field in fields):
        raise SystemExit("session receipt does not match manifest")
    executable_hash = sha256_file(args.executable.resolve())
    if receipt.get("executable_sha256") != executable_hash: raise SystemExit("receipt executable hash mismatch")
    if not isinstance(receipt.get("executable_path"), str) or Path(receipt["executable_path"]).resolve() != args.executable.resolve():
        raise SystemExit("receipt executable path mismatch")
    if (receipt.get("complete") is not True or not isinstance(receipt.get("git_commit"), str) or
        not re.fullmatch(r"[0-9a-fA-F]{40}", receipt["git_commit"]) or
        not isinstance(receipt.get("dirty"), bool) or
        not isinstance(receipt.get("feature_configuration"), str) or
        not isinstance(receipt.get("start_time_unix"), int) or
        not isinstance(receipt.get("end_time_unix"), int) or
        receipt["end_time_unix"] < receipt["start_time_unix"]):
        raise SystemExit("receipt is incomplete")
    entry = verified_entry(args.root, manifest["variant_id"])
    if entry["content_id"] != manifest["content_id"]: raise SystemExit("inventory content identity changed")
    data = verdicts(args.root)
    for prior in data["verdicts"].values():
        for mode_record in prior.get("modes", {}).values():
            if mode_record.get("nonce") == manifest["nonce"]: raise SystemExit("session receipt nonce was already collected")
    record = data["verdicts"].setdefault(entry["variant_id"], {
        "variant_id": entry["variant_id"], "content_id": entry["content_id"],
        "track_id": entry["track_id"], "version": entry["version"],
        "lev_media_id": entry["lev_media_id"], "vrm_media_id": entry["vrm_media_id"],
        "lev_sha256": manifest["lev_sha256"], "vrm_sha256": manifest["vrm_sha256"], "modes": {}})
    record["modes"][manifest["mode"]] = {"result": args.result, "notes": args.notes or "",
        "launch_method": "user-steam-attested", "reviewer": args.reviewer,
        "git_commit": receipt["git_commit"], "dirty": bool(receipt.get("dirty")),
        "feature_configuration": receipt["feature_configuration"],
        "start_time_unix": receipt["start_time_unix"], "end_time_unix": receipt["end_time_unix"],
        "executable_sha256": executable_hash, "nonce": manifest["nonce"],
        "recorded_at": datetime.now(timezone.utc).isoformat()}
    atomic_json(args.root / "verdicts.json", data)
    consumed = args.root / "session" / f"manifest.{manifest['nonce']}.consumed.json"
    os.replace(manifest_path, consumed)
    os.replace(receipt_path, args.root / "session" / f"receipt.{manifest['nonce']}.consumed.json")
    print(f"Collected Steam-attested {manifest['mode']}={args.result} for {entry['variant_id']}.")


def command_copy_verdict(args: argparse.Namespace) -> None:
    source = verified_entry(args.root, args.source_variant_id)
    destination = verified_entry(args.root, args.destination_variant_id)
    if source["variant_id"] == destination["variant_id"]: raise SystemExit("source and destination variants must differ")
    if source["content_id"] != destination["content_id"]: raise SystemExit("variants do not share exact content bytes")
    for entry in (source, destination):
        report = entry["verification"]
        for role in ("lev", "vrm"):
            path = (args.root / entry[f"{role}_path"]).resolve()
            if not path.is_file() or path.is_symlink() or sha256_file(path) != report[f"{role}_sha256"]:
                raise SystemExit(f"{entry['variant_id']} {role.upper()} no longer matches inventory")
    data = verdicts(args.root)
    source_record = data["verdicts"].get(source["variant_id"])
    if not source_record or not source_record.get("modes"): raise SystemExit("source variant has no verdicts to copy")
    if destination["variant_id"] in data["verdicts"]: raise SystemExit("destination already has a verdict record")
    copied_at = datetime.now(timezone.utc).isoformat()
    modes = {}
    for mode, record in source_record["modes"].items():
        modes[mode] = dict(record, copied_from_variant=source["variant_id"], copied_by=args.reviewer,
                           copied_at=copied_at, copy_reason=args.reason)
    data["verdicts"][destination["variant_id"]] = {
        "variant_id": destination["variant_id"], "content_id": destination["content_id"],
        "track_id": destination["track_id"], "version": destination["version"],
        "lev_media_id": destination["lev_media_id"], "vrm_media_id": destination["vrm_media_id"],
        "lev_sha256": destination["verification"]["lev_sha256"],
        "vrm_sha256": destination["verification"]["vrm_sha256"], "modes": modes}
    atomic_json(args.root / "verdicts.json", data)
    print(f"Copied {len(modes)} explicit verdict(s) from {source['variant_id']} to {destination['variant_id']}.")


def command_abort_session(args: argparse.Namespace) -> None:
    session = args.root / "session"
    source = session / "armed.json"
    if not source.exists(): source = session / "manifest.json"
    manifest = load_json(source)
    if not isinstance(manifest, dict) or not isinstance(manifest.get("nonce"), str):
        raise SystemExit("no prepared session to abort")
    os.replace(source, session / f"manifest.{manifest['nonce']}.aborted.json")
    receipt = session / "receipt.json"
    if receipt.exists(): os.replace(receipt, session / f"receipt.{manifest['nonce']}.aborted.json")
    print(f"Aborted and consumed review session {manifest['nonce']}.")


def rows_for_export(root: Path) -> list[dict[str, Any]]:
    ledger = verdicts(root)["verdicts"]
    rows = []
    for entry in inventory(root)["entries"]:
        row = {key: entry.get(key) for key in ("track_id", "name", "author", "track_type", "collection",
                                                "version", "lev_media_id", "vrm_media_id", "lev_crc32",
                                                "vrm_crc32", "status", "notes", "variant_id", "content_id")}
        report = entry.get("verification", {})
        row.update({"lev_sha256": report.get("lev_sha256"), "vrm_sha256": report.get("vrm_sha256"),
                    "existing_saphi_tags": ", ".join(entry.get("saphi_modes", [])),
                    "runtime_support": "Unsupported" if entry.get("track_type") == "battle" else "Steam review"})
        runtime = ledger.get(entry.get("variant_id", ""), {}).get("modes", {})
        differences = []
        for mode in MODES:
            analysis = report.get("file_analysis", {}).get(mode, {}).get("result", "Error")
            tested = runtime.get(mode, {}).get("result", "Not tested")
            row[f"{mode} File analysis"] = analysis
            row[f"{mode} In-game test"] = tested
            tagged = mode in entry.get("saphi_modes", [])
            if analysis == "Detected" and not tagged: differences.append(f"{mode}: structure detected, Saphi absent")
            if analysis == "Not detected" and tagged: differences.append(f"{mode}: structure not detected, Saphi present")
        row["differences"] = "; ".join(differences)
        rows.append(row)
    return rows


def command_export(args: argparse.Namespace) -> None:
    rows = rows_for_export(args.root)
    output = {"schema_version": SCHEMA_VERSION, "generated_at": datetime.now(timezone.utc).isoformat(),
              "source": TRACKS_URL, "modes": list(MODES), "entries": rows}
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
    except ImportError as exc:
        raise SystemExit("openpyxl is required for spreadsheet export") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Verification"
    headers = list(rows[0]) if rows else ["track_id"]
    sheet.append(headers)
    escaped = False
    for row in rows:
        sheet.append([row.get(header) for header in headers])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str) and cell.value.lstrip(" \t")[:1] in {"=", "+", "-", "@"}:
                cell.value = "'" + cell.value
                cell.number_format = "@"
                escaped = True
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = min(55, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[column[0].column_letter].width = width
    if rows:
        table = Table(displayName="SaphiVerification", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    generation_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S.%fZ-") + secrets.token_hex(8)
    exports = args.root / "exports"
    exports.mkdir(parents=True, exist_ok=True)
    staging = Path(tempfile.mkdtemp(prefix=f".{generation_id}.", dir=exports))
    report_json = staging / "report.json"
    report_xlsx = staging / "report.xlsx"
    try:
        atomic_json(report_json, output)
        workbook.save(report_xlsx)
        with report_xlsx.open("rb") as stream: os.fsync(stream.fileno())
        manifest = {"schema_version": SCHEMA_VERSION, "generation_id": generation_id,
                    "row_count": len(rows), "xlsx_text_escaped": escaped,
                    "report_json_sha256": sha256_file(report_json),
                    "report_xlsx_sha256": sha256_file(report_xlsx)}
        atomic_json(staging / "manifest.json", manifest)
        generation = exports / generation_id
        os.replace(staging, generation)
        directory = os.open(exports, os.O_RDONLY)
        try: os.fsync(directory)
        finally: os.close(directory)
        atomic_json(exports / "current.json", {"schema_version": SCHEMA_VERSION,
                                                "generation_id": generation_id})
    except BaseException:
        if staging.exists(): shutil.rmtree(staging)
        raise
    print(f"published {len(rows)} rows as export generation {generation_id}")


def command_cleanup_exports(args: argparse.Namespace) -> None:
    exports = args.root / "exports"
    current = load_json(exports / "current.json")
    if not isinstance(current, dict) or not isinstance(current.get("generation_id"), str):
        raise SystemExit("no valid current export generation")
    generations = sorted(path for path in exports.iterdir() if path.is_dir() and not path.name.startswith("."))
    removable = [path for path in generations if path.name != current["generation_id"]]
    removed = 0
    for path in removable[:-args.keep] if args.keep else removable:
        manifest = load_json(path / "manifest.json")
        if (not isinstance(manifest, dict) or manifest.get("generation_id") != path.name or
            sha256_file(path / "report.json") != manifest.get("report_json_sha256") or
            sha256_file(path / "report.xlsx") != manifest.get("report_xlsx_sha256")):
            raise SystemExit(f"refusing to remove unverified export generation {path.name}")
        shutil.rmtree(path)
        removed += 1
    print(f"removed {removed} non-current verified export generation(s)")


def parser() -> argparse.ArgumentParser:
    result = argparse.ArgumentParser()
    result.add_argument("--root", type=Path, default=Path("custom-content-lab"))
    commands = result.add_subparsers(dest="command", required=True)
    sync = commands.add_parser("sync")
    sync.add_argument("--catalogue-file", type=Path)
    sync.set_defaults(func=command_sync)
    analyze = commands.add_parser("analyze")
    analyze.add_argument("--verifier", type=Path, required=True)
    analyze.add_argument("--offline", action="store_true")
    analyze.set_defaults(func=command_analyze)
    queue = commands.add_parser("queue")
    queue.add_argument("--difference", action="store_true")
    queue.add_argument("--author")
    queue.add_argument("--track-type", choices=("race", "battle"))
    queue.add_argument("--mode", choices=MODES)
    queue.add_argument("--collection")
    queue.add_argument("--pending", action="store_true")
    queue.add_argument("--newest-only", action="store_true")
    queue.add_argument("--page-size", type=int, default=50, choices=range(1, 201))
    queue.add_argument("--cursor", type=int, default=0)
    queue.add_argument("--json", action="store_true")
    queue.set_defaults(func=command_queue)
    prepare = commands.add_parser("prepare-session")
    prepare.add_argument("variant_id")
    prepare.add_argument("mode", choices=MODES)
    prepare.add_argument("--host", type=int, required=True)
    prepare.add_argument("--minutes", type=int, default=30)
    prepare.set_defaults(func=command_prepare_session)
    collect = commands.add_parser("collect-session")
    collect.add_argument("--executable", type=Path, required=True)
    collect.add_argument("--reviewer", required=True)
    collect.add_argument("--steam-attested", action="store_true")
    collect.add_argument("--result", choices=("Yes", "No"), required=True)
    collect.add_argument("--notes")
    collect.set_defaults(func=command_collect_session)
    abort = commands.add_parser("abort-session")
    abort.set_defaults(func=command_abort_session)
    copy = commands.add_parser("copy-verdict")
    copy.add_argument("source_variant_id")
    copy.add_argument("destination_variant_id")
    copy.add_argument("--reviewer", required=True)
    copy.add_argument("--reason", required=True)
    copy.set_defaults(func=command_copy_verdict)
    export = commands.add_parser("export")
    export.set_defaults(func=command_export)
    cleanup = commands.add_parser("cleanup-exports")
    cleanup.add_argument("--keep", type=int, default=2, choices=range(0, 101))
    cleanup.set_defaults(func=command_cleanup_exports)
    return result


def main() -> int:
    args = parser().parse_args()
    args.root = args.root.resolve()
    args.func(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
